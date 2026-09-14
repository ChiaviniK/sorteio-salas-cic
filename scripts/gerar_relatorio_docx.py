# -*- coding: utf-8 -*-
"""Gera docs/RELATORIO_TECNICO.docx a partir de docs/RELATORIO_TECNICO.md.

Conversor Markdown -> Word minimalista (títulos, tabelas, listas, citações,
blocos de código e negrito/itálico/código inline), com estilo institucional Unesp.
Uso: python scripts/gerar_relatorio_docx.py
"""
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

RAIZ = Path(__file__).resolve().parent.parent
ORIGEM = RAIZ / "docs" / "RELATORIO_TECNICO.md"
DESTINO = RAIZ / "docs" / "RELATORIO_TECNICO.docx"

AZUL_UNESP = RGBColor(0x00, 0x35, 0x94)
AZUL_ESCURO = RGBColor(0x00, 0x27, 0x76)
CINZA = RGBColor(0x47, 0x55, 0x69)

RE_INLINE = re.compile(r"(\*\*.+?\*\*|\*[^*]+?\*|`[^`]+`)")


def preencher_runs(paragrafo, texto):
    """Aplica **negrito**, *itálico* e `código` inline."""
    for parte in RE_INLINE.split(texto):
        if not parte:
            continue
        if parte.startswith("**") and parte.endswith("**"):
            run = paragrafo.add_run(parte[2:-2])
            run.bold = True
        elif parte.startswith("*") and parte.endswith("*") and len(parte) > 2:
            run = paragrafo.add_run(parte[1:-1])
            run.italic = True
        elif parte.startswith("`") and parte.endswith("`"):
            run = paragrafo.add_run(parte[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
        else:
            paragrafo.add_run(parte)


def sombrear_celula(celula, cor_hex):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), cor_hex)
    celula._tc.get_or_add_tcPr().append(shd)


def adicionar_rodape_com_pagina(doc):
    secao = doc.sections[0]
    par = secao.footer.paragraphs[0]
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = par.add_run("Sorteio de Salas · CIC Unesp — Relatório Técnico · página ")
    run.font.size = Pt(8)
    run.font.color.rgb = CINZA
    campo_inicio = OxmlElement("w:fldChar")
    campo_inicio.set(qn("w:fldCharType"), "begin")
    instrucao = OxmlElement("w:instrText")
    instrucao.text = "PAGE"
    campo_fim = OxmlElement("w:fldChar")
    campo_fim.set(qn("w:fldCharType"), "end")
    run_pagina = par.add_run()
    run_pagina.font.size = Pt(8)
    run_pagina.font.color.rgb = CINZA
    run_pagina._r.append(campo_inicio)
    run_pagina._r.append(instrucao)
    run_pagina._r.append(campo_fim)


def montar():
    linhas = ORIGEM.read_text(encoding="utf-8").splitlines()
    doc = Document()

    secao = doc.sections[0]
    secao.page_width = Cm(21.0)
    secao.page_height = Cm(29.7)
    secao.left_margin = secao.right_margin = Cm(2.5)
    secao.top_margin = secao.bottom_margin = Cm(2.2)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)

    for nivel, tamanho in (("Heading 1", 15), ("Heading 2", 12.5)):
        estilo = doc.styles[nivel]
        estilo.font.name = "Calibri Light"
        estilo.font.size = Pt(tamanho)
        estilo.font.bold = True
        estilo.font.color.rgb = AZUL_UNESP if nivel == "Heading 1" else AZUL_ESCURO

    adicionar_rodape_com_pagina(doc)

    i = 0
    while i < len(linhas):
        linha = linhas[i]

        if linha.startswith("# "):
            titulo = doc.add_paragraph()
            run = titulo.add_run(linha[2:].strip())
            run.font.name = "Calibri Light"
            run.font.size = Pt(21)
            run.bold = True
            run.font.color.rgb = AZUL_UNESP
            titulo.paragraph_format.space_after = Pt(14)
            i += 1
            continue

        if linha.startswith("### "):
            doc.add_heading(linha[4:].strip(), level=2)
            i += 1
            continue

        if linha.startswith("## "):
            doc.add_heading(linha[3:].strip(), level=1)
            i += 1
            continue

        if linha.strip() in ("---", ""):
            i += 1
            continue

        if linha.startswith("|"):
            bloco = []
            while i < len(linhas) and linhas[i].startswith("|"):
                celulas = [c.strip() for c in linhas[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r"-{3,}", c) for c in celulas):
                    bloco.append(celulas)
                i += 1
            colunas = max(len(l) for l in bloco)
            tabela = doc.add_table(rows=len(bloco), cols=colunas)
            tabela.style = "Table Grid"
            tabela.alignment = WD_TABLE_ALIGNMENT.CENTER
            for r, valores in enumerate(bloco):
                for c in range(colunas):
                    celula = tabela.cell(r, c)
                    celula.paragraphs[0].paragraph_format.space_after = Pt(2)
                    preencher_runs(celula.paragraphs[0], valores[c] if c < len(valores) else "")
                    for par in celula.paragraphs:
                        for run in par.runs:
                            run.font.size = Pt(9.5)
                            if r == 0:
                                run.bold = True
                    if r == 0:
                        sombrear_celula(celula, "DDE8FA")
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
            continue

        if linha.startswith("```"):
            i += 1
            codigo = []
            while i < len(linhas) and not linhas[i].startswith("```"):
                codigo.append(linhas[i])
                i += 1
            i += 1
            for linha_codigo in codigo:
                par = doc.add_paragraph()
                par.paragraph_format.left_indent = Cm(0.75)
                par.paragraph_format.space_after = Pt(0)
                run = par.add_run(linha_codigo if linha_codigo else " ")
                run.font.name = "Consolas"
                run.font.size = Pt(9)
                run.font.color.rgb = CINZA
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
            continue

        if linha.startswith("> "):
            par = doc.add_paragraph()
            par.paragraph_format.left_indent = Cm(1.0)
            par.paragraph_format.space_after = Pt(4)
            run_base = par.add_run("")
            run_base.font.name = "Cambria"
            preencher_runs(par, linha[2:].strip())
            for run in par.runs:
                run.font.name = "Cambria"
                run.font.size = Pt(10.5)
                run.font.color.rgb = AZUL_ESCURO
            i += 1
            continue

        if linha.startswith("- "):
            item = []
            texto = linha[2:]
            j = i + 1
            while j < len(linhas) and linhas[j].startswith("  ") and not linhas[j].startswith("- "):
                texto += " " + linhas[j].strip()
                j += 1
            par = doc.add_paragraph(style="List Bullet")
            preencher_runs(par, texto.strip())
            i = j
            continue

        num = re.match(r"^(\d+)\.\s+(.*)$", linha)
        if num:
            texto = num.group(2)
            j = i + 1
            while j < len(linhas) and linhas[j].startswith("   ") and linhas[j].strip():
                texto += " " + linhas[j].strip()
                j += 1
            par = doc.add_paragraph(style="List Number")
            preencher_runs(par, texto.strip())
            i = j
            continue

        # parágrafo corrido (agrega linhas contíguas)
        texto = linha.strip()
        j = i + 1
        while j < len(linhas) and linhas[j].strip() and not re.match(
            r"^(#|\||```|- |> |---|\d+\.\s)", linhas[j]
        ):
            texto += " " + linhas[j].strip()
            j += 1
        par = doc.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        preencher_runs(par, texto)
        i = j

    doc.save(DESTINO)
    print(f"Gerado: {DESTINO}")


if __name__ == "__main__":
    montar()
