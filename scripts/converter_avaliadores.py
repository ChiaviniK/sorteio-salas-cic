# -*- coding: utf-8 -*-
"""Converte a lista de avaliadores (docx) para o template xlsx do app."""
import re
import unicodedata

from docx import Document
from openpyxl import Workbook

BASE = r"C:\Users\lenovo\Desktop\projetoscopilot\cic\dados-exemplo"

doc = Document(BASE + r"\Avaliadores CIC.docx")
linhas = [p.text.strip() for p in doc.paragraphs if p.text.strip()]


def norm(s: str) -> str:
    s = unicodedata.normalize("NFD", s)
    return "".join(c for c in s if unicodedata.category(c) != "Mn").lower()


wb = Workbook()
ws = wb.active
ws.title = "Avaliadores"
ws.append(["Nome", "E-mail", "Função", "Tipo", "Instituição/Unidade"])

FUNCOES = {"docente": "Docente", "pos": "Pós-graduando", "apta": "Externo", "convidado": "Externo"}
INSTITUICOES = {"docente": "Unesp", "pos": "Unesp", "apta": "Apta", "convidado": "Convidado"}

secao = "docente"
contagem: dict[str, int] = {}
for linha in linhas:
    low = norm(linha)
    if "doutorando" in low or "graduando" in low:
        secao = "pos"
        continue
    if low == "apta":
        secao = "apta"
        continue
    if low == "convidados":
        secao = "convidado"
        continue
    if re.match(r"^\d+\s", linha):
        break
    nome = linha
    tipo = "Geral"
    if re.search(r"pibic\s*jr", low):
        tipo = "PIBIC Jr"
        nome = re.sub(r"\s*[–\-\u2013\u2010\u2011]\s*pibic\s*jr\.?\s*$", "", nome, flags=re.I)
        nome = re.sub(r"\s*\(pibic\s*jr\.?\)", "", nome, flags=re.I).strip()
    funcao = FUNCOES[secao]
    contagem[funcao] = contagem.get(funcao, 0) + 1
    ws.append([nome, "", funcao, tipo, INSTITUICOES[secao]])

for coluna, largura in zip("ABCDE", [40, 30, 18, 12, 30]):
    ws.column_dimensions[coluna].width = largura

wb.save(BASE + r"\Avaliadores_CIC.xlsx")
print("Total:", ws.max_row - 1, "|", contagem)
